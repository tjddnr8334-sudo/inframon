#!/usr/bin/env python3
"""7p 자료의 **말만** 다듬는다 — 뜻은 그대로, 쓰는 말을 우리 분야 말로.

용어 하나씩 고칠 때마다 스크립트를 새로 만들면 따라가기 어려워진다. 이 파일에 모아
두고 필요할 때 항목만 덧붙인다. 어느 것을 왜 바꿨는지는 아래 목록과 로그에 남는다.

    python scripts/patch_deck_wording.py [--dry-run]
    python scripts/rebuild_7p_vF.py          # 이어서 vF 판을 다시 만든다

산출: 원본 7p 자료를 제자리에서 갱신 · docs/bridges/deck_wording_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from patch_deck_soften_7p import cell_text                            # noqa: E402
from patch_kaia_v2_slide13 import set_para                            # noqa: E402

DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
LOG: list = []

# 표 칸: (바꾸기 전, 바꾼 뒤, 이유) — **글자로 찾는다**.
#
# 칸 번호로 찾으면 손으로 고친 사본(로고를 넣거나 도형을 옮긴 판)에서 자리가 밀려
# 엉뚱한 칸을 덮어쓴다. 글자로 찾으면 어느 판에든 같은 곳에 꽂힌다 — `--deck` 으로
# 대상만 바꿔 주면 된다.
CELLS = [
    ("실측 건강 교량으로 기준 재설정 → 잠정 해제",
     "안전등급 양호 교량의 실측값으로 기준 확정",
     "'건강 교량' 은 healthy bridge 직역이다 — 국내 유지관리 말은 **안전등급**(A~E)이라 "
     "'안전등급 양호' 로 바꾼다. 그리고 '기준 재설정 → 잠정 해제' 는 같은 일을 두 번 "
     "말한 것이다(기준을 확정하면 잠정이 풀린다). 한 번만 적고 '확정' 으로 끝낸다 — "
     "같은 행 '지금 단계' 칸의 '(기준 잠정)' 과 짝이 맞는다"),
]

# 본문 도형: (슬라이드, 도형, 문단, 바꾸기 전 글의 일부, [강조 묶음별 글], 이유)
EDITS: list = []


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--log", default="docs/bridges/deck_wording_log.json")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    ok = True

    for needle, new, why in CELLS:
        hits = []
        for sn, sl in enumerate(prs.slides, start=1):
            for sh in sl.shapes:
                if not sh.has_table:
                    continue
                for ri, row in enumerate(sh.table.rows):
                    for ci, c in enumerate(row.cells):
                        if needle in c.text:
                            hits.append((sn, ri, ci, c))
        if not hits:
            LOG.append({"찾음": False, "찾던 글": needle})
            ok = False
            continue
        for sn, ri, ci, c in hits:
            old = c.text if a.dry_run else cell_text(c, new)
            LOG.append({"쪽": sn, "자리": f"표 r{ri}c{ci}", "이전": old, "이후": new,
                        "글자수": f"{len(old)} → {len(new)}", "이유": why})

    for sn, shi, pi, needle, parts, why in EDITS:
        sh = prs.slides[sn - 1].shapes[shi]
        pa = sh.text_frame.paragraphs[pi]
        cur = "".join(r.text for r in pa.runs)
        if needle not in cur:
            LOG.append({"찾음": False, "쪽": sn, "자리": f"[{shi}]p{pi}",
                        "찾던 글": needle, "그 자리 글": cur[:80]})
            ok = False
            continue
        old = cur if a.dry_run else set_para(sh, pi, parts)
        LOG.append({"쪽": sn, "자리": f"[{shi}]p{pi}", "이전": old,
                    "이후": "".join(parts), "이유": why})

    for e in LOG:
        if e.get("찾음") is False:
            print(f"  ⚠ 못 찾음: {e['찾던 글']}")
        else:
            print(f"  · {e['쪽']}쪽 {e['자리']:<10} {e['이전']}")
            print(f"       → {e['이후']}   ({e.get('글자수', '')})")
    if not ok:
        print("!! 자리가 안 맞는다 — 아무것도 저장하지 않는다")
        return 1
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name,
         "_왜": "뜻은 그대로 두고 쓰는 말을 우리 분야 말로 바꾼다.",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {a.deck}")
    print("wrote", a.log)
    print("   다음: python scripts/rebuild_7p_vF.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
