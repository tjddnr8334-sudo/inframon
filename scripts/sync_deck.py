#!/usr/bin/env python3
"""손으로 고친 7p 사본을 **지금까지의 글 수정에 따라잡게** 한다.

작업 파일이 갈라졌다. 원본(`7p_1`)과 생성본(`7p_vF`) 말고도 손으로 고친 사본이
여럿 생겼고(_로고수정 · _찐찐 · _찐찐_song), 그 사본들에는 그동안의 글 수정이 들어가
있지 않다. 사본마다 도형 번호가 밀려 있으므로 번호로 찾으면 엉뚱한 곳을 덮어쓴다.

그래서 여태 만든 패치 스크립트들의 **수정 목록만 모아** 글자로 찾아 넣는다. 이미
반영된 항목은 찾을 글이 없으므로 조용히 건너뛴다 — **몇 번을 돌려도 같은 결과**가
되므로, 사본이 또 생겨도 이 한 줄이면 따라잡는다.

    python scripts/sync_deck.py --deck "docs/..._찐찐_song.pptx" [--dry-run]

산출: 대상 pptx 를 제자리에서 갱신 · docs/bridges/deck_sync_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

import patch_deck_kict_asks as kict                                   # noqa: E402
import patch_deck_sidecar as sidecar                                  # noqa: E402
import patch_deck_text as textp                                       # noqa: E402
import patch_deck_wording as wording                                  # noqa: E402
from patch_kaia_v2_slide13 import set_para, style_groups              # noqa: E402


def collect() -> list:
    """패치 스크립트들의 수정 목록을 (찾을 글, [새 글 조각], 출처) 로 모은다.

    스크립트마다 자리를 가리키는 방식이 다르다(도형 번호 · 표 칸 번호 · 글자). 여기서는
    **자리 정보를 버리고 글자만** 쓴다 — 사본에서는 자리가 못 미덥기 때문이다.
    순서는 원래 목록 순서를 지킨다. 번호를 올리는 수정(⑤→⑥ 먼저)처럼 순서가
    중요한 것이 있다.
    """
    out = []
    for e in textp.EDITS:                       # (찾을 글, [조각], 이유)
        out.append((e[0], e[1], "text", None))
    for e in wording.CELLS:                     # (찾을 글, 새 글, 이유)
        out.append((e[0], [e[1]], "wording", None))
    for e in sidecar.EDITS:                     # (쪽, 도형, 문단, 찾을 글, [조각], 이유)
        out.append((e[3], e[4], "sidecar", e[0]))
    for e in sidecar.CELLS:                     # (쪽, 행, 열, 찾을 글, 새 글, 이유)
        out.append((e[3], [e[4]], "sidecar", e[0]))
    for e in kict.EDITS:                        # (도형, 문단, 찾을 글, 새 글, 이유)
        out.append((e[2], [e[3]], "kict", 7))   # 이 목록은 전부 Ⅳ장(7쪽)이다
    return out


# 문단 일부만 바꾸는 것: (찾을 글, 그 자리에 넣을 글, 쪽)
#
# 목차 줄처럼 **긴 문장 안의 한 토막**만 고쳐야 하는 경우가 있다. 문단을 통째로
# 갈아 끼우면 나머지 글이 날아간다 — dry-run 에서 1쪽 목차가 통째로 지워질 뻔했다.
SUBSTR = [
    ("KICT 협의 요청", "KICT 자료 요청", 1),
]

# 자료에 남겨 둔 메모 — 그 용어를 고치고 나면 같이 치운다.
NOTES = [
    "건강 교량?? ->용어가 어색함",          # → '안전등급 양호 교량' 으로 고쳤다
    "사람 손 없이 -> 표현 어색함",          # → '자동으로' 로 고쳤다
    "사이드카 라는 의미가",                 # → '조회 전용 프로그램' 으로 풀었다
    "0에서 5번 X",                         # → 단계를 ① ~ ⑥ 으로 고쳤다
    "좌표하나로 끝까지라는 표어",             # → 다섯 군데 중 두 곳을 덜어냈다
    "벌써 23개소 완료",                     # → '1 개 · 넣는 것은 좌표뿐' 으로
    "벌써 16개소 완료",                     # → '12 일 · 위성이 다시 찍는 주기' 로
    "12개가 무엇을 뜻하는지",                # → '0 개 · 교량에 다는 장비' 로
    "3D로 나온게 맞는지",                   # → 그림 설명에 '(높이는 추정값)' 을 넣었다
]


def find(prs, needle, slide=None):
    hits = []
    for sn, sl in enumerate(prs.slides, start=1):
        if slide is not None and sn != slide:
            continue
        for shi, sh in enumerate(sl.shapes):
            holders = []
            if sh.has_text_frame:
                holders.append((sh, f"[{shi}]"))
            elif sh.has_table:
                for ri, row in enumerate(sh.table.rows):
                    for ci, c in enumerate(row.cells):
                        holders.append((c, f"표 r{ri}c{ci}"))
            for holder, where in holders:
                for pi, pa in enumerate(holder.text_frame.paragraphs):
                    if needle in "".join(r.text for r in pa.runs):
                        hits.append((sn, where, pi, holder))
    return hits


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", required=True)
    ap.add_argument("--log", default="docs/bridges/deck_sync_log.json")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    log, done, skip, amb = [], 0, 0, 0

    for needle, parts, src, slide in collect():
        hits = find(prs, needle, slide)
        if not hits:
            skip += 1
            continue                            # 이미 반영됐다 — 조용히 넘어간다
        if len(hits) > 1:
            amb += 1
            log.append({"여러 곳": needle, "출처": src,
                        "자리": [f"{s}쪽 {w}p{p}" for s, w, p, _ in hits]})
            print(f"  ⚠ 여러 곳에 있어 건드리지 않는다: {needle}")
            continue
        sn, where, pi, holder = hits[0]
        old = "".join(r.text for r in holder.text_frame.paragraphs[pi].runs)
        if not a.dry_run:
            set_para(holder, pi, parts)
        done += 1
        log.append({"쪽": sn, "자리": f"{where}p{pi}", "출처": src,
                    "이전": old, "이후": "".join(parts)})
        print(f"  · {sn}쪽 {where}p{pi:<2} {old[:38]}")
        print(f"       → {''.join(parts)[:38]}")

    # 문단 일부만 바꾸는 것 — 나머지 글은 그대로 둔다.
    for needle, repl, slide in SUBSTR:
        for sn, where, pi, holder in find(prs, needle, slide):
            pa = holder.text_frame.paragraphs[pi]
            old = "".join(r.text for r in pa.runs)
            if len(style_groups(pa.runs)) > 1:
                print(f"  ⚠ 겉모습이 여러 갈래라 건드리지 않는다: {old[:40]}")
                continue
            if not a.dry_run:
                set_para(holder, pi, [old.replace(needle, repl)])
            done += 1
            log.append({"쪽": sn, "자리": f"{where}p{pi}", "출처": "일부",
                        "이전": old, "이후": old.replace(needle, repl)})
            print(f"  · {sn}쪽 {where}p{pi:<2} {old[:38]}")
            print(f"       → {old.replace(needle, repl)[:38]}")

    # 자료에 남겨 둔 메모 치우기
    for memo in NOTES:
        for sl in prs.slides:
            for sh in list(sl.shapes):
                if sh.has_text_frame and memo in sh.text_frame.text:
                    if not a.dry_run:
                        sh._element.getparent().remove(sh._element)
                    done += 1
                    log.append({"메모 치움": sh.text_frame.text[:40]})
                    print(f"  − 메모 치움: «{memo}»")

    print(f"바꿈 {done} · 이미 반영돼 건너뜀 {skip}"
          + (f" · 모호해서 건드리지 않음 {amb}" if amb else ""))
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0
    if not done:
        print("바꿀 것이 없다 — 이미 최신이다")
        return 0

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name,
         "_왜": "손으로 고친 사본은 도형 번호가 밀려 있어 글자로 찾아 넣는다. 이미 반영된 "
              "항목은 건너뛰므로 몇 번을 돌려도 같은 결과가 된다.",
         "바꿈": done, "건너뜀": skip, "바꾼 것": log},
        ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {a.deck}")
    print("wrote", a.log)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
