#!/usr/bin/env python3
"""7p 자료의 글을 **글자로 찾아** 바꾼다 — 어느 판에든 같은 곳에 꽂히도록.

손으로 고친 사본이 여러 개 생겼다(_로고수정 · _찐찐 · _찐찐_song). 도형 번호로 찾으면
로고를 넣거나 도형을 옮긴 판에서 자리가 밀려 엉뚱한 곳을 덮어쓴다. 그래서 **바꾸기 전
글자**로 찾는다. `--deck` 으로 대상만 바꿔 주면 어느 판에든 같은 수정이 들어간다.

이번에 고치는 것

  ① 4쪽 단계 번호가 ⓪ 부터 시작했다 → ① ~ ⑥ 으로. 사람은 1부터 센다.
  ② 4쪽 수치 타일 — '51 장면 · 23 개소 · 16 개소 · 12 개' 는 무엇을 뜻하는지 알기
     어렵고, 전부 "벌써 이만큼 했다" 로 읽힌다. **무엇을 어떻게 하는지**로 바꾼다
     (6단계 · 좌표 1개 · 12일 · 장비 0개).
  ③ '좌표 하나로 끝까지' 가 다섯 군데 반복됐다 → 표지와 Ⅱ장 카드, 4쪽 타일에만 남긴다.
  ④ 5쪽 '사람 손 없이' → '자동으로'. 같은 뜻인데 뒤쪽이 덜 튄다.
  ⑤ 5쪽 그림 설명에 **높이는 추정값**임을 적는다. 점 높이는 DEM + 잔차고도 Δh 이고
     Δh 의 불확실도가 교량마다 7.6~22 m(중앙 13.5 m)다. 주탑처럼 30~50 m 차이는
     가르지만 교각·교대는 못 가린다 — 안 적어 두면 그 자리에서 질문이 나온다.
  ⑥ 7쪽 향후 계획에 PINN 이 없었다 → 성능 평가 · 노후도 추정 · 거동 예측을 넣는다.

    python scripts/patch_deck_text.py [--deck <파일>] [--dry-run]

산출: 대상 pptx 를 제자리에서 갱신 · docs/bridges/deck_text_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from patch_kaia_v2_slide13 import set_para                            # noqa: E402

DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
LOG: list = []

# (바꾸기 전 글자, [겉모습 묶음별 새 글], 이유)
# **순서가 중요하다** — 앞의 수정이 만든 글자를 뒤의 찾기가 잡지 않도록 짠다.
EDITS = [
    # ── 4쪽 단계 번호: 뒤에서부터 올린다(⑤→⑥ 먼저라야 안 겹친다) ──────────
    ("⑤ 산출 · 연동", ["⑥ 산출 · 연동"], "사람은 1부터 센다 — ⓪ 로 시작하지 않는다"),
    ("④ 위험도 평가", ["⑤ 위험도 평가"], "위와 같음"),
    ("③ PINN 역산", ["④ PINN 역산"], "위와 같음"),
    ("② 교면 측점 정합", ["③ 교면 측점 정합"], "위와 같음"),
    ("① InSAR 처리", ["② InSAR 처리"], "위와 같음"),
    ("⓪ 교량·영상 선별", ["① 교량·영상 선별"], "위와 같음"),

    # ── 4쪽 수치 타일: 뒤에서부터(12→0 먼저라야 16→12 와 안 겹친다) ────────
    ("12  개", ["0", "  개"], "'12개 엔드포인트' 는 무엇을 뜻하는지 알기 어렵다"),
    ("B-Maps REST 엔드포인트", ["교량에 다는 장비"], "위 타일에 맞춤"),
    ("16  개소", ["12", "  일"], "'16개소 완료' 로 읽힌다 — 갱신 주기로 바꾼다"),
    ("한강 교량 트윈 시범", ["위성이 다시 찍는 주기"], "위 타일에 맞춤"),
    ("23  개소", ["1", "  개"], "'23개소 완료' 로 읽힌다 — 넣는 것이 무엇인지로"),
    ("전 과정 시범 처리 교량", ["넣는 것은 좌표뿐"], "위 타일에 맞춤"),
    ("51  장면", ["6", "  단계"], "'51장면' 은 규모지 이해를 돕지 않는다 — 파이프라인 길이로"),
    ("실제 Sentinel-1 영상", ["영상 고르기부터 B-Maps 전달까지"], "위 타일에 맞춤"),

    # ── 되풀이되는 표어 덜어내기 ──────────────────────────────────────────
    ("Inframon 은 좌표 하나로 위성",
     ["위성 영상에서 변위를 뽑아 AI 로 해석하고, 트윈으로 만들어 B-Maps 에 "
      "넘기는 6단계다"],
     "'좌표 하나로' 가 다섯 군데 반복됐다 — 여기서는 무엇을 하는지만"),
    ("좌표 1개 실행", ["같은 절차로"], "같은 이유"),
    ("교량을 바꿔도 코드 수정이 없다", ["교량이 바뀌어도 고칠 것이 없다"],
     "'코드' 는 듣는 분들의 말이 아니다"),
    ("좌표 1개로 수집–처리–교면 변위 시계열까지 자동",
     ["수집–처리–교면 변위 시계열까지 자동"], "표어 반복을 덜어낸다"),

    # ── 5쪽 ──────────────────────────────────────────────────────────────
    ("교량마다 사람 손 없이",
     ["교량마다 같은 산출물과 IFC 트윈이 자동으로 나온다"],
     "'사람 손 없이' 보다 '자동으로' 가 덜 튄다. 되풀이되던 표어도 함께 덜어낸다"),
    ("IFC 트윈 + 교면 위성 측점 (산출물 원본)",
     ["IFC 트윈 + 교면 위성 측점 (높이는 추정값)"],
     "점 높이는 DEM + 잔차고도 Δh 이고 Δh 불확실도가 7.6~22 m 다. 주탑은 가르지만 "
     "교각·교대는 못 가린다 — 안 적어 두면 그 자리에서 질문이 나온다"),

    # ── 7쪽 향후 계획에 PINN 을 넣는다 ───────────────────────────────────
    ("계측 위치 기반 정확도 검증 · CRI 기준 재설정",
     ["PINN 성능 평가 · 노후도 추정 · 거동 예측"],
     "향후 계획에 PINN 이 한 줄도 없었다 — 이 엔진이 앞으로 무엇을 하는지 적는다"),
    ("코너리플렉터 · 고해상도 SAR 보강",
     ["코너리플렉터로 측점 확보 · 정확도 검증"],
     "보강 그 자체보다 그것으로 무엇을 하는지 — 정확도 검증을 여기로 옮긴다"),
]


def find(prs, needle):
    """그 글자가 들어 있는 문단을 전부 찾는다 → [(쪽, 도형, 문단번호, 도형객체)]."""
    hits = []
    for sn, sl in enumerate(prs.slides, start=1):
        for shi, sh in enumerate(sl.shapes):
            frames = []
            if sh.has_text_frame:
                frames.append((sh, f"[{shi}]"))
            elif sh.has_table:
                for ri, row in enumerate(sh.table.rows):
                    for ci, c in enumerate(row.cells):
                        frames.append((c, f"표 r{ri}c{ci}"))
            for holder, where in frames:
                for pi, pa in enumerate(holder.text_frame.paragraphs):
                    if needle in "".join(r.text for r in pa.runs):
                        hits.append((sn, where, pi, holder))
    return hits


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--log", default="docs/bridges/deck_text_log.json")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    ok = True

    for needle, parts, why in EDITS:
        hits = find(prs, needle)
        if not hits:
            LOG.append({"찾음": False, "찾던 글": needle})
            ok = False
            continue
        if len(hits) > 1:
            LOG.append({"찾음": "여러 곳", "찾던 글": needle,
                        "자리": [f"{s}쪽 {w}p{p}" for s, w, p, _ in hits]})
            ok = False
            continue
        sn, where, pi, holder = hits[0]
        old = "".join(r.text for r in holder.text_frame.paragraphs[pi].runs)
        if not a.dry_run:
            set_para(holder, pi, parts)
        LOG.append({"쪽": sn, "자리": f"{where}p{pi}", "이전": old,
                    "이후": "".join(parts), "이유": why})

    for e in LOG:
        if e.get("찾음") is False:
            print(f"  ⚠ 못 찾음: {e['찾던 글']}")
        elif e.get("찾음") == "여러 곳":
            print(f"  ⚠ 여러 곳에 있다: {e['찾던 글']} → {e['자리']}")
        else:
            print(f"  · {e['쪽']}쪽 {e['자리']:<12} {e['이전'][:40]}")
            print(f"       → {e['이후'][:40]}")
    if not ok:
        print("!! 자리가 안 맞는다 — 아무것도 저장하지 않는다")
        return 1
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name,
         "_왜": "단계 번호를 1부터, 알기 어려운 수치 타일을 '무엇을 어떻게 하는지' 로, "
              "되풀이되던 표어를 덜어내고, 향후 계획에 PINN 을 넣는다.",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {a.deck}")
    print("wrote", a.log)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
