#!/usr/bin/env python3
"""4쪽 하단 오른쪽 — 수치 타일 넷을 **원리상 못 하는 것** 세 줄로 바꾼다.

타일 넷(6단계 · 1개 · 12일 · 0개)은 전부 다른 쪽에 또 있다. 6단계는 바로 위
파이프라인 띠가 보여 주고, 12일·0개는 2쪽 큰 숫자에 있고, '좌표 하나' 는 왼쪽 칸과
3쪽 카드에 있다. 같은 말을 네 번 더 하는 자리다.

그 자리에 **우리가 못 하는 것**을 적는다. 각 교량 `결과.md` 의 "이 파이프라인이
원리상 못 하는 것" 에서 가져온다. 4쪽이 "이렇게 돌아갑니다 + 이건 못 합니다" 가 되어
균형이 잡히고, 자문위원이 물어볼 것을 **우리 입으로 먼저** 답하는 자리가 된다.

자료 전체가 "성능은 두루뭉실, 목적은 분명" 으로 가는 중이다. 한계를 먼저 말하는
것만큼 그 톤에 맞는 것이 없다.

    python scripts/patch_deck_p4_limits.py [--deck <파일>] [--dry-run]

산출: 대상 pptx 를 제자리에서 갱신
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from kaia_theme import BLUE, GRAY, INK, PAPER, SLATE, box, panel, text  # noqa: E402

DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"

# 지울 타일 — 글자로 찾는다(사본마다 도형 번호가 다르다).
TILE_TEXTS = [
    "6  단계", "영상 고르기부터 B-Maps 전달까지",
    "1  개", "넣는 것은 좌표뿐",
    "12  일", "위성이 다시 찍는 주기",
    "0  개", "교량에 다는 장비",
]

# (제목, 왜 못 하는지)
LIMITS = [
    ("교면 측점 밀도",
     "위성 화소가 지상 5×20 m 다. 데크 폭은 1~2 화소라 프로그램으로 늘릴 수 없다"),
    ("절대 강성(EI)",
     "상대 변위만 보므로 관측으로 낼 수 없다 — 설계 제원으로 계산한 값이다"),
    ("연직 변위",
     "한 궤도로는 위성 시선 방향만 나온다. 두 궤도를 합쳐야 위아래가 된다"),
]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--slide", type=int, default=4)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    s = prs.slides[a.slide - 1]

    # ── 타일 지우기 ───────────────────────────────────────────────────────
    victims, xs, ys = [], [], []
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip() in TILE_TEXTS:
            victims.append(sh)
            xs += [sh.left, sh.left + sh.width]
            ys += [sh.top, sh.top + sh.height]
    if not victims:
        print("!! 지울 타일을 못 찾았다 — 이미 바꿨거나 다른 자료다")
        return 1
    # 글자 없는 배경 상자도 같이 — 타일이 차지하던 네모 안에 들어 있는 것만.
    x0, x1, y0, y1 = min(xs), max(xs), min(ys), max(ys)
    for sh in list(s.shapes):
        if sh in victims or sh.left is None:
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            continue
        if (x0 - 228600 <= sh.left and sh.left + sh.width <= x1 + 228600
                and y0 - 228600 <= sh.top and sh.top + sh.height <= y1 + 228600):
            victims.append(sh)

    IN = 914400
    print(f"  지울 도형 {len(victims)}개 · 자리 "
          f"x{x0/IN:.2f}~{x1/IN:.2f} y{y0/IN:.2f}~{y1/IN:.2f}")
    if a.dry_run:
        for v in victims:
            t = v.text_frame.text.strip() if v.has_text_frame else "(배경)"
            print(f"     − «{t[:30]}»")
        print("(dry-run — 저장하지 않았다)")
        return 0
    for v in victims:
        v._element.getparent().remove(v._element)

    # ── 그 자리에 '원리상 못 하는 것' ─────────────────────────────────────
    px, pw = x0 / IN - 0.2, (x1 - x0) / IN + 0.4
    py, ph = 4.35, 2.25
    panel(s, px, py, pw, ph, "원리상 못 하는 것", fill=PAPER, tab=SLATE)
    yy = py + 0.46
    for i, (head, why) in enumerate(LIMITS):
        box(s, px + 0.18, yy + 0.06, 0.05, 0.2, fill=BLUE)
        text(s, px + 0.34, yy, pw - 0.52, 0.24, [(head, 11, True, INK)])
        text(s, px + 0.34, yy + 0.24, pw - 0.52, 0.42, [(why, 9.2, False, GRAY)])
        yy += 0.60 if i < len(LIMITS) - 1 else 0
    prs.save(a.deck)
    print(f"고쳤다: {a.deck}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
