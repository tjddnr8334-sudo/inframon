#!/usr/bin/env python3
"""7p 자료에서 **'사이드카' 를 걷어낸다** — 듣는 사람이 모르는 말이다.

사이드카는 오토바이 옆칸에서 온 비유로, 본체를 고치지 않고 옆에 따로 띄우는 프로그램을
가리킨다. 쿠버네티스·마이크로서비스 쪽에서는 흔한 말이지만, 구조·유지관리 하시는
분들에게는 처음 듣는 단어다. 자문회의에서 'FRAM 공명' 이 어렵다는 말이 나온 것과
같은 종류의 문제다.

말을 바꾸면서 **자리를 옮긴다.** 이 칸들에서 KICT 담당자가 판단하는 것은 "우리가 뭘
해야 하나" 이고, 거기서 가장 듣고 싶은 말은 **"손댈 게 없습니다"** 다. 그래서
'사이드카' 가 있던 자리에 '본체 수정 없이' 와 '조회 전용' 을 넣는다.

REST 는 다 없애지 않는다 — 6쪽 연동 접점표의 `REST /api/v1` 은 기술 규격을 적는
칸이라 그대로 둔다. 기술 담당자가 "어떤 방식이냐" 물을 때 답이 되는 말이다.

    python scripts/patch_deck_sidecar.py [--dry-run]
    python scripts/rebuild_7p_vF.py          # 이어서 vF 판을 다시 만든다

산출: 원본 7p 자료를 제자리에서 갱신 · docs/bridges/deck_sidecar_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from patch_deck_soften_7p import cell_text                            # noqa: E402
from patch_kaia_v2_slide13 import set_para                            # noqa: E402

DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
LOG: list = []

# (슬라이드, 도형, 문단, 바꾸기 전 글의 일부, [강조 묶음별 글], 이유)
EDITS = [
    # ── Ⅱ 연구성과 목표 ④번 카드 ─────────────────────────────────────────
    (3, 42, 0, "REST 사이드카로", ["본체 수정 없이"],
     "이 칸에서 가장 듣고 싶은 말은 '손댈 게 없습니다' 다 — 그 말을 앞에 세운다"),
    (3, 42, 1, "B-Maps 탭에 결과 제공", ["조회 전용 서버가 탭에 결과 제공"],
     "무엇이 결과를 주는지 한 말로 — '조회 전용' 이 사이드카를 대신한다"),

    # ── Ⅲ③ B-Maps 연동 ──────────────────────────────────────────────────
    (6, 14, 0, "읽기 전용 사이드카로 붙인다",
     ["B-Maps 본체를 고치지 않고 옆에 붙는 조회 전용 프로그램으로 잇는다 — "
      "교량관리번호로 조회하면 교면 변위와 IFC 트윈이 그대로 올라온다"],
     "머리글에서 비유를 빼고 무엇인지 그대로 적는다"),
    (6, 26, 1, "사이드카 · 읽기 전용", ["조회 전용 서버"],
     "도해 아래 설명에 이미 '본체 수정 없음' 이 있으므로 겹치지 않게 짧게"),
    (6, 44, 0, "읽기 전용 사이드카라",
     ["※ 교면 전용 재선별(MT-InSAR: 기선망 · 점별 DEM 오차 · 열팽창 · APS) 뒤의 "
      "교면 PS 를 쓴다. 연동은 조회 전용이라 B-Maps 본체를 고치지 않는다."],
     "각주도 같은 말로"),

    # ── Ⅳ 향후 계획 ─────────────────────────────────────────────────────
    (7, 23, 2, "B-Maps REST 사이드카 · 탭 예제", ["B-Maps 조회 API · 탭 예제"],
     "지나온 일을 적는 칸 — 무엇을 만들었는지만"),
]

# 표 칸: (슬라이드, 행, 열, 바꾸기 전, 바꾼 뒤, 이유)
CELLS = [
    (3, 4, 2, "REST 사이드카 · 탭 예제 구현", "조회 전용 API · 탭 예제 구현",
     "성과 지표표의 '지금 단계' 칸 — 위 카드와 말을 맞춘다"),
]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--log", default="docs/bridges/deck_sidecar_log.json")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    ok = True

    for sn, shi, pi, needle, parts, why in EDITS:
        sh = prs.slides[sn - 1].shapes[shi]
        pa = sh.text_frame.paragraphs[pi]
        cur = "".join(r.text for r in pa.runs)
        if needle not in cur:
            LOG.append({"찾음": False, "쪽": sn, "도형": shi, "문단": pi,
                        "찾던 글": needle, "그 자리 글": cur[:80]})
            ok = False
            continue
        old = cur if a.dry_run else set_para(sh, pi, parts)
        LOG.append({"쪽": sn, "자리": f"[{shi}]p{pi}", "이전": old,
                    "이후": "".join(parts), "이유": why})

    for sn, ri, ci, needle, new, why in CELLS:
        tb = next(s for s in prs.slides[sn - 1].shapes if s.has_table).table
        cur = tb.cell(ri, ci).text
        if needle not in cur:
            LOG.append({"찾음": False, "쪽": sn, "자리": f"표 r{ri}c{ci}",
                        "찾던 글": needle, "그 자리 글": cur[:80]})
            ok = False
            continue
        old = cur if a.dry_run else cell_text(tb.cell(ri, ci), new)
        LOG.append({"쪽": sn, "자리": f"표 r{ri}c{ci}", "이전": old,
                    "이후": new, "이유": why})

    for e in LOG:
        if e.get("찾음") is False:
            print(f"  ⚠ {e['쪽']}쪽 못 찾음: {e['찾던 글']} / 그 자리: {e['그 자리 글']}")
        else:
            print(f"  · {e['쪽']}쪽 {e['자리']:<10} {e['이전'][:46]}")
            print(f"       → {e['이후'][:46]}")
    if not ok:
        print("!! 자리가 안 맞는다 — 아무것도 저장하지 않는다")
        return 1
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name,
         "_왜": "'사이드카' 는 오토바이 옆칸에서 온 비유로, 본체를 고치지 않고 옆에 띄우는 "
              "프로그램을 뜻한다. 구조·유지관리 하시는 분들에게는 처음 듣는 말이라 "
              "'본체 수정 없이' · '조회 전용' 으로 바꾼다.",
         "_남긴것": "6쪽 연동 접점표의 REST /api/v1 — 기술 규격을 적는 칸이라 그대로 둔다",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {a.deck}")
    print("wrote", a.log)
    print("   다음: python scripts/rebuild_7p_vF.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
