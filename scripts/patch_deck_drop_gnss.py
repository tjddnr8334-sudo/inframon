#!/usr/bin/env python3
"""발표자료에서 **계측 대조·품질 판정·집계 수치를 걷어낸다**.

`strip_gnss_from_bridges.py` 가 교량 결과물에서 한 일을 발표자료에도 한다. 이유는
같다 — 지금 값으로는 계측 대조가 아무것도 주장하지 못하는데, 그 표를 붙여 두면
읽는 사람에게 남는 건 "이 데이터는 계측과 안 맞는다" 한 줄뿐이다. 아직 완성 전이라
성능으로 다툴 자리가 아니다. 축은 **왜 결합해야 하는가**로 되돌린다.

걷어내는 것
  · 슬라이드 3 표 — '현장 보고서 판정 대조 / 16개소 중 14개소 일치' 행 통째
    '23개소(한강 16개소 포함)' · '한강 16개소 산출' · '16/16개소 · 결합률 94%'
  · 슬라이드 5 — 표의 **품질 판정 열**(조건부 · 전송 보류) · '16/16개소' · '94% 결합률'
    제목의 '16개소 전수 처리' · 각주의 결합률 분수
  · 슬라이드 6 — 오른쪽 '현장 계측 대조' 열 통째(1/5 · 4/5 개소 · R² · 우연 기준선)
    → 계측을 빼고 **교면 전용 MT-InSAR 산출**로 바꾼다
  · 표지 목차의 '현장 검증'

남기는 것: 코너리플렉터·고해상도 SAR 제안(Ⅳ). 계측 대조가 아니라 **측점 밀도**가
한계라는 말로 바꿔 적는다. 향후계획(Ⅳ)의 '계측 위치 기반 정확도 검증' 은 앞으로
하겠다는 계획이지 지금의 주장이 아니라 그대로 둔다.

    python scripts/patch_deck_drop_gnss.py

산출: 대상 pptx 를 제자리에서 갱신 · docs/bridges/deck_drop_gnss_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from patch_research_deck import find, set_text                        # noqa: E402

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
LOG: list = []


def swap(slide, needle, runs, why):
    """도형 하나의 글자를 바꾼다 — 있는 run 만 건드린다(구조를 깨지 않는다)."""
    sh = find(slide, needle)
    if sh is None:
        LOG.append({"찾음": False, "찾던 글": needle})
        return
    old = set_text(sh, runs)
    LOG.append({"이전": old.replace(chr(10), " / ")[:130],
                "이후": " / ".join("".join(t for t, _ in pr) for pr in runs)[:130],
                "이유": why})


def cell_text(cell, text):
    """표 칸의 첫 run 글자만 바꾼다 — 서식을 지킨다."""
    pa = cell.text_frame.paragraphs[0]
    if not pa.runs:
        return
    pa.runs[0].text = text
    for r in pa.runs[1:]:
        r.text = ""


def drop_row(table, idx, why):
    LOG.append({"표 행 삭제": " | ".join(c.text for c in table.rows[idx].cells)[:110],
                "이유": why})
    tbl = table._tbl
    tbl.remove(tbl.tr_lst[idx])


def drop_col(table, idx, why):
    """열 하나를 빼고 그 폭을 남은 열에 비례해 나눠 준다."""
    LOG.append({"표 열 삭제": table.cell(0, idx).text, "이유": why})
    gone = table.columns[idx].width
    tbl = table._tbl
    grid = tbl.find(A_NS + "tblGrid")
    grid.remove(list(grid)[idx])
    for tr in tbl.tr_lst:
        tcs = [c for c in tr if c.tag == A_NS + "tc"]
        tr.remove(tcs[idx])
    rest = [c.width for c in table.columns]
    tot = sum(rest)
    for c, w in zip(table.columns, rest):
        c.width = w + int(round(gone * w / tot))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--log", default="docs/bridges/deck_drop_gnss_log.json")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    S = prs.slides

    # ── 표지 목차 ─────────────────────────────────────────────────────────
    sh = find(S[0], "B-Maps 연동 · 현장 검증")
    if sh is not None:
        for pa in sh.text_frame.paragraphs:
            for r in pa.runs:
                if "현장 검증" in r.text:
                    r.text = r.text.replace("현장 검증", "서비스 제공")
        LOG.append({"이전": "③ B-Maps 연동 · 현장 검증",
                    "이후": "③ B-Maps 연동 · 서비스 제공",
                    "이유": "검증을 주장할 자리가 아니다 — 연동과 제공이 이 장의 내용이다"})

    # ── 슬라이드 3 · 성과 지표 표 ─────────────────────────────────────────
    tb3 = next(s for s in S[2].shapes if s.has_table).table
    drop_row(tb3, 2,
             "'현장 보고서 판정 대조 / 16개소 중 14개소 일치' — 이 값은 데크 ±30 m 로 고른 "
             "강변 지반 점에서 나왔고, 교면 전용으로 다시 재도 지금 값으로는 일치를 "
             "주장하지 못한다")
    for ri, txt, why in (
        (1, "좌표 1개로 수집–처리–교면 변위 시계열까지 자동",
         "'23개소(한강 16개소 포함)' — 개소 수는 달성이 아니라 규모다. 무엇이 되는지로 바꿈"),
        (2, "열·하중 성분 분리 · CRI 산출 (기준 잠정)",
         "'한강 16개소 산출' 에서 개소 수를 뺌"),
        (3, "IFC 부재 GUID 에 교면 측점 결합",
         "'16/16개소 · 결합률 94%' — 비율은 검증 성적처럼 읽힌다"),
    ):
        cell_text(tb3.cell(ri, 2), txt)
        LOG.append({"이전": f"슬라이드3 표 r{ri} 현재 달성", "이후": txt, "이유": why})
    swap(S[2], "판정 대조는 2024",
         [[("※ 현재 달성 = 실제 Sentinel-1 자료(ASC path127, 2018-06-19 ~ 2025-12-27) "
            "처리 결과 기준.", None)]],
         "각주에서 '판정 대조' 출처 문장을 뺌 — 대조를 싣지 않으므로")

    # ── 슬라이드 5 · 한강 교량 실증 ───────────────────────────────────────
    swap(S[4], "16개소 전수 처리",
         [[("교량마다 사람 손 없이 같은 산출물과 IFC 트윈이 나온다 — 좌표 하나로 끝까지",
            None)]],
         "'16개소 전수 처리' 는 개소 수 자랑이다. 자동화가 요지다")
    tb5 = next(s for s in S[4].shapes if s.has_table).table
    drop_col(tb5, 5, "'품질 판정'(조건부 · 전송 보류) — 게이트 낱말은 구조 등급으로 오독된다")
    swap(S[4], "품질 판정 ≠ 구조 등급",
         [[("연 변위속도 = 교면 측점 중앙값(mm/년) · 관측 시점 수가 다르면 신뢰구간도 다르다",
            None)]],
         "'현장 GNSS 계측 7개소' 와 '품질 판정' 을 뺌")
    swap(S[4], "한강 교량 16개소 IFC 트윈",
         [[("IFC 트윈 + 교면 위성 측점 (산출물 원본)", None)]], "개소 수를 뺌")
    swap(S[4], "16/16", [[("IFC", None), ("  4x3", None)]],
         "'16/16 개소' — 전수 달성 비율을 산출물 규격으로 바꿈")
    swap(S[4], "IFC 트윈 생성", [[("부재 GUID 포함 트윈", None)]], "위 타일에 맞춤")
    swap(S[4], "94", [[("3D", None), ("  Tiles", None)]],
         "'94 % 결합률' — 비율은 검증 성적처럼 읽힌다")
    swap(S[4], "측점–부재 결합률", [[("지도 레이어 출력", None)]], "결합률 타일을 산출물로 바꿈")
    swap(S[4], "결합률 = 1,113",
         [[("※ 표의 값은 각 교량 폴더의 bridge.json · 결과.md · project.h5 에서 직접 읽음.",
            None)]],
         "각주에서 결합률 분수를 뺌")

    # ── 슬라이드 6 · B-Maps 연동 ──────────────────────────────────────────
    swap(S[5], "연구수행 내용 ③", [[("연구수행 내용 ③ — B-Maps 연동", None)]],
         "'현장 검증' 을 제목에서 뺌")
    swap(S[5], "5개소 중 1개소만 일치",
         [[("B-Maps 본체를 고치지 않고 읽기 전용 사이드카로 붙인다 — 교량관리번호로 "
            "조회하면 교면 변위와 IFC 트윈이 그대로 올라온다", None)]],
         "머리글에서 대조 결과를 빼고 연동이 무엇인지로 바꿈")
    swap(S[5], "현장 계측 대조", [[("교면 전용 MT-InSAR 산출", "FFFFFF")]],
         "오른쪽 열을 계측 대조에서 처리 산출로 바꿈")
    swap(S[5], "1/5", [[("4종", None), ("  값", None)]], "'1/5 개소 일치' 를 뺌")
    swap(S[5], "교면 PS 가 계측과 일치",
         [[("측점마다 산출", None)], [("속도 · 잔차고도 · 열팽창 · 코히런스", None)]],
         "R² · 우연 기준선을 빼고 점별 산출값으로")
    swap(S[5], "4/5", [[("APS", None), ("  제거", None)]], "'4/5 개소 미도출' 을 뺌")
    swap(S[5], "계측 지점에서 PS 도출 못함",
         [[("기선망 뒤 동시 추정", None)],
          [("수직·시간 기선망 · 점별 DEM 오차", None)]],
         "계측 지점 PS 개수를 빼고 처리 내용으로")
    swap(S[5], "원인  ",
         [[("한계  ", "ED6C00"),
           ("교면 측점 밀도가 한계다 — Sentinel-1 IW 는 지상 5×20 m 이고 데크 폭은 "
            "1~2 화소다. 주경간에 코너리플렉터를 놓거나 고해상도 SAR 을 쓰면 교면 "
            "측점이 채워진다 (Ⅳ 참조)", None)]],
         "한계의 근거를 계측 대조가 아니라 분해능으로 적음 — 코너리플렉터 제안은 남긴다")
    swap(S[5], "※ 대조는 교면 전용",
         [[("※ 교면 전용 재선별(MT-InSAR: 기선망 · 점별 DEM 오차 · 열팽창 · APS) 뒤의 "
            "교면 PS 를 쓴다. 연동은 읽기 전용 사이드카라 B-Maps 본체를 고치지 않는다.",
            None)]],
         "각주에서 대조 방법 설명을 빼고 산출·연동 방식으로")

    swap(S[4], "출처·판정 문서", [[("결과 기록", None)], [("출처·근거 문서", None)]],
         "'판정 문서' 의 판정 낱말을 뺌")

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name,
         "_왜": "지금 값으로는 계측 대조가 아무것도 주장하지 못한다. 성능으로 다툴 자리가 "
              "아니므로 축을 '왜 결합해야 하는가' 로 되돌린다.",
         "_짝": "docs/bridges 쪽은 scripts/strip_gnss_from_bridges.py 가 같은 일을 한다",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print("고쳤다:", a.deck)
    for e in LOG:
        if "이전" in e:
            print(f"  · {e['이전'][:44]:<46} → {e['이후'][:44]}")
        elif e.get("찾음") is False:
            print("  ⚠ 못 찾음:", e["찾던 글"])
        else:
            print("  −", list(e.values())[0][:70])
    print("wrote", a.log)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
