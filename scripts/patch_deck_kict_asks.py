#!/usr/bin/env python3
"""7p 자료 Ⅳ장 — **KICT 에게 일을 나누자고 하지 않는다. 자료를 청한다.**

지금 네 칸은 시범 교량 선정 · 화면 필드 매핑 · 운영 규칙 합의 · 검증 자료 공유였다.
넷 다 **이미 같이 하기로 한 사이에서 나올 말**이다. 자문회의에 오시는 분들은
Inframon 을 보러 오는 것이고, 합치자고 말을 꺼내는 쪽은 우리다. 그 자리에서 상대에게
과제를 배정하면 순서가 뒤집힌다.

그래서 네 칸을 전부 **B-Maps 가 이미 들고 있는 자료를 주십사** 하는 청으로 바꾼다.
받을 만한 것만 적는다 — 달라고 해서 이상하지 않고, 받으면 바로 결과가 좋아지는 것.

  · 교량 대장 제원 — 지금은 공개데이터와 OSM 으로 **추정**한다. 실제로 어긋난다:
    등록 연장과 맞는 OSM way 가 없어 시점–종점 선분을 쓴 교량이 여럿이고, 폭은
    보도 간격으로 짐작했다(올림픽대교는 그 값이 106 m 로 나왔다). **교면 측점을
    가려내는 회랑 폭이 이 제원에서 나온다** — 제원이 맞으면 선별이 바로 좋아진다.
  · 교량관리번호와 대장 좌표 — 결과를 대장에 붙이는 열쇠다. 지금은 이름으로 맞춘다.
  · 보유 IFC 모델 — 우리는 공개데이터로 세운 IFC proxy 를 쓴다. 실 모델이 오면
    우리 트윈이 아니라 **B-Maps 트윈의 부재**에 값이 붙는다.
  · 준공·보수 이력 — 변위 시계열에 계단이 보일 때 보수인지 이상인지 가른다.
    연도와 일자만 있으면 된다.

    python scripts/patch_deck_kict_asks.py [--dry-run]
    python scripts/rebuild_7p_vF.py          # 이어서 vF 판을 다시 만든다

산출: 원본 7p 자료를 제자리에서 갱신 · docs/bridges/deck_kict_asks_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from patch_kaia_v2_slide13 import set_para, style_groups               # noqa: E402

DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
LOG: list = []

# (도형, 문단, 바꾸기 전 글의 일부, 바꾼 뒤, 이유)
EDITS = [
    (38, 0, "KICT 협의 요청", "KICT 자료 요청",
     "과제를 배정하는 말투에서 자료를 청하는 말투로"),

    (41, 0, "시범 교량 선정", "교량 대장 제원",
     "받을 만한 것 — 지금은 공개데이터·OSM 추정이라 등록 제원과 어긋나는 교량이 있다"),
    (41, 1, "계측 자료가 있는 곳부터", "연장 · 폭 · 경간 · 형하고",
     "교면 측점을 가려내는 회랑 폭이 이 값에서 나온다"),

    (44, 0, "화면 필드 매핑", "교량관리번호",
     "화면 규격을 같이 정하자는 말 대신, 결과를 대장에 붙일 열쇠를 청한다"),
    (44, 1, "교량관리번호 · B-Maps 표시 항목", "대장 좌표와 함께",
     "지금은 교량 이름으로 맞추고 있다"),

    (47, 0, "운영 규칙 합의", "보유 IFC 모델",
     "운영 규칙은 연동이 정해진 뒤의 이야기다 — 지금 받을 수 있는 것으로"),
    (47, 1, "갱신 주기 · 경보 임계 · 인증", "있는 교량만이라도",
     "우리는 공개데이터로 세운 IFC proxy 를 쓴다. 실 모델이 오면 B-Maps 부재에 바로 붙는다"),

    (50, 0, "검증 자료 공유", "준공 · 보수 이력",
     "'검증'은 같이 하기로 한 뒤의 말이다. 대장에 있는 이력만 청한다"),
    (50, 1, "계측 위치 · 실 제원 · 점검 이력", "연도와 일자만",
     "변위 시계열의 계단이 보수인지 이상인지 가르는 데 쓴다"),
]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--slide", type=int, default=7)
    ap.add_argument("--log", default="docs/bridges/deck_kict_asks_log.json")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    s = prs.slides[a.slide - 1]
    ok = True

    for shi, pi, needle, new, why in EDITS:
        sh = s.shapes[shi]
        pa = sh.text_frame.paragraphs[pi]
        cur = "".join(r.text for r in pa.runs)
        if needle not in cur:
            LOG.append({"찾음": False, "도형": shi, "문단": pi,
                        "찾던 글": needle, "그 자리 글": cur[:80]})
            ok = False
            continue
        g = style_groups(pa.runs)
        size = (None if pa.runs[g[0][0]].font.size is None
                else pa.runs[g[0][0]].font.size.pt)
        old = cur if a.dry_run else set_para(sh, pi, [new])
        LOG.append({"도형": shi, "문단": pi, "이전": old, "이후": new,
                    "이유": why, "크기pt": size, "글자수": f"{len(cur)} → {len(new)}"})

    for e in LOG:
        if e.get("찾음") is False:
            print(f"  ⚠ [{e['도형']}]p{e['문단']} 못 찾음: {e['찾던 글']}"
                  f" / 그 자리: {e['그 자리 글']}")
        else:
            print(f"  · [{e['도형']}]p{e['문단']} {e['크기pt']}pt  "
                  f"{e['이전']:<26} → {e['이후']}   ({e['글자수']}자)")
    if not ok:
        print("!! 자리가 안 맞는다 — 아무것도 저장하지 않는다")
        return 1
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name, "_슬라이드": a.slide,
         "_왜": "자문회의에 오시는 분들은 Inframon 을 보러 오는 것이고, 합치자고 말을 "
              "꺼내는 쪽은 우리다. 그 자리에서 상대에게 과제를 배정하면 순서가 뒤집힌다. "
              "네 칸을 전부 B-Maps 가 이미 들고 있는 자료를 주십사 하는 청으로 바꾼다.",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {a.deck}")
    print("wrote", a.log)
    print("   다음: python scripts/rebuild_7p_vF.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
