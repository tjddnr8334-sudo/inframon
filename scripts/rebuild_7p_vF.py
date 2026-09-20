#!/usr/bin/env python3
"""7p 자료를 **`2026 KAIA 자문회의 ppt_vF.pptx` 와 같은 형식**으로 다시 짠다.

색과 글꼴은 이미 같았다 — `kaia_theme.py` 의 10171E · 373D47 · 1C5B8A · Paperlogy 는
vF 에서 뽑아 온 값이다. 다른 것은 **틀**이었다. vF 는 머리띠·왼쪽 세로띠·쪽번호·
바닥띠·연구센터명을 슬라이드마다 그리지 않고 **레이아웃(`page 1`)** 에 두고, 슬라이드는
장 번호 배지와 기관 로고, 소제목 칩만 얹는다. 7p 자료는 그 장식을 쪽마다 직접
그리고 있었고(도형 0~18), 로고도 없었다.

그래서 vF 에서 **틀만 뽑아** 그 위에 7p 의 내용을 옮긴다.

  ① `--make-template` — vF 에서 슬라이드 18장을 떼고 레이아웃·마스터·테마만 남긴다.
     끼워넣은 글꼴 21개(46 MB)와 과한 해상도의 표지 그림(5760×3240, 11 MB)은 덜어내
     284 MB → 1.9 MB 로 만든다. Paperlogy 는 이 컴퓨터에 깔려 있고 7p 자료도 원래
     끼워넣지 않고 썼다.
  ② 표지는 vF 표지 형식으로 다시 그린다 — 짙은 띠(053B79) · 흰 세로바 · 기관 로고
     여섯 개 · 아래 소관부처 줄. 목차는 그 아래에 넣어 **7쪽을 유지한다**
     (vF 는 표지와 목차를 두 쪽으로 나누지만, 이 자료는 7쪽이 이름이다).
  ③ 본문 6쪽은 `page 1` 레이아웃에 얹고, 원본의 **내용 도형만** 옮긴다.
     원본 도형 0~12(머리띠·장 번호·Ⅰ~Ⅳ 탭)와 15~18(바닥띠·쪽번호)은 레이아웃이
     대신하므로 버리고, 13~14(소제목 줄)와 19 이후를 그대로 가져온다.

그림은 관계(rId)를 새 슬라이드에 다시 맺어 준다 — 그러지 않으면 복사된 `<p:pic>` 이
원본 슬라이드의 rId 를 가리켜 빈 칸으로 열린다.

    python scripts/rebuild_7p_vF.py --make-template      # 틀을 새로 뽑을 때만
    python scripts/rebuild_7p_vF.py

산출: docs/KICT_Bmaps_Inframon_연구개발_7p_vF.pptx (원본은 건드리지 않는다)
"""

from __future__ import annotations

import argparse
import copy
import io
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from kaia_theme import (                                              # noqa: E402
    BLUE, C, FONT_B, GRAY, PAPER, RULE, SLATE, WHITE, _typeface,
    box, emu, text,
)

VF = ROOT / "docs" / "2026 KAIA 자문회의 ppt_vF.pptx"
TEMPLATE = ROOT / "docs" / "templates" / "kaia_vF_page.pptx"
SRC = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
OUT = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_vF.pptx"
LOGO = ROOT / "docs" / "img" / "kaia_logo"

SW = 13.333
NAVY = C("053B79")          # vF 표지 띠
CHROME = set(range(0, 13)) | {15, 16, 17, 18}   # 레이아웃이 대신하는 원본 장식

# 본문 쪽: (원본 쪽 번호, 장 번호 배지, 레이아웃 제목)
PAGES = [
    (2, "01", "연구개발과제의 필요성"),
    (3, "02", "연구성과 목표"),
    (4, "03", "연구수행 내용 ① — Inframon 플랫폼 구성"),
    (5, "03", "연구수행 내용 ② — 한강 교량 실증 · 디지털 트윈"),
    (6, "03", "연구수행 내용 ③ — B-Maps 연동"),
    (7, "04", "향후 계획"),
]

# vF 표지의 기관 로고 — 자리도 vF 와 같게 둔다.
COVER_LOGOS = [("s1_14.jpg", 0.09, 6.70, 1.34, 0.43),
               ("s1_10.png", 1.64, 6.73, 1.97, 0.35),
               ("s1_7.png", 3.82, 6.73, 1.91, 0.41),
               ("s1_8.png", 5.94, 6.75, 2.08, 0.40),
               ("s1_9.jpg", 8.23, 6.60, 2.19, 0.56),
               ("s1_13.png", 10.63, 6.82, 2.41, 0.32)]

AGENCIES = ("[소관부처] 국토교통부 │ [전문기관] 국토교통과학기술진흥원 │ "
            "[주관연구개발기관] 강원대학교 │ [공동연구개발기관] 스마트인사이드에이아이 │ "
            "[자문/협력기관] 한국건설기술연구원(KICT), 지티에스엔지니어링(GTS)")

CONTENTS = [
    ("Ⅰ", "연구개발과제의 필요성",
     "교량 유지관리 현황과 한계 · 위성 InSAR 와 B-Maps 연계의 필요성", "2"),
    ("Ⅱ", "연구성과 목표", "최종 목표 · 세부 목표 · 성과 지표와 지금 단계", "3"),
    ("Ⅲ", "연구수행 내용",
     "① Inframon 플랫폼 구성   ② 한강 교량 실증 · 디지털 트윈   ③ B-Maps 연동", "4 – 6"),
    ("Ⅳ", "향후 계획", "단계별 추진 계획 · KICT 자료 요청 · 기대 효과", "7"),
]


# ── ① vF → 틀 ─────────────────────────────────────────────────────────────
def make_template(vf: Path, out: Path) -> None:
    """vF 에서 슬라이드·끼워넣은 글꼴·과한 표지 그림을 덜어 틀만 남긴다."""
    from PIL import Image

    tmp = out.parent / (out.stem + "_full.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(vf, tmp)

    prs = Presentation(tmp)
    lst = prs.slides._sldIdLst
    n = len(list(lst))
    for sldId in list(lst):
        prs.part.drop_rel(sldId.get(qn("r:id")))
        lst.remove(sldId)
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    for el in prs.part._element.findall(P + "embeddedFontLst"):
        prs.part._element.remove(el)
    prs.save(tmp)

    zin = zipfile.ZipFile(tmp)
    prels = re.sub(r'<Relationship[^>]*Target="fonts/[^"]+"[^>]*/>', "",
                   zin.read("ppt/_rels/presentation.xml.rels").decode("utf-8"))
    ct = zin.read("[Content_Types].xml").decode("utf-8").replace(
        '<Default Extension="fntdata" ContentType="application/x-fontdata"/>', "")
    im = Image.open(io.BytesIO(zin.read("ppt/media/image1.png")))
    im.thumbnail((1800, 1800), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in zin.infolist():
            if item.filename.startswith("ppt/fonts/"):
                continue
            data = {"ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "ppt/media/image1.png": buf.getvalue()}.get(
                        item.filename, zin.read(item.filename))
            z.writestr(item, data)
    zin.close()
    tmp.unlink()
    print(f"틀: {out}  ({os.path.getsize(vf) / 1e6:.0f} MB → "
          f"{os.path.getsize(out) / 1e6:.2f} MB · 슬라이드 {n}장 제거)")


# ── ② 도형 옮기기 ─────────────────────────────────────────────────────────
def copy_shape(dst, src_slide, shape, seed: int) -> None:
    """도형 하나를 XML 째로 옮기고, 그림이면 관계를 새로 맺는다."""
    el = copy.deepcopy(shape._element)
    for blip in el.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        # 원본 파트를 그대로 참조하면 **이름이 충돌한다** — 원본의 image1.png 와
        # 틀의 image1.png 가 같은 자리를 다퉈 zip 에 같은 이름이 두 번 들어간다.
        # 바이트만 가져와 새 패키지에 새 이름으로 넣는다.
        blob = src_slide.part.related_part(rid).blob
        _, new_rid = dst.part.get_or_add_image_part(io.BytesIO(blob))
        blip.set(qn("r:embed"), new_rid)
    # 도형 id 는 슬라이드 안에서 겹치면 안 된다 — 새 번호를 준다.
    for i, nv in enumerate(el.iter(qn("p:cNvPr"))):
        nv.set("id", str(seed + i))
    dst.shapes._spTree.append(el)


def chrome(slide, tag: str, title: str) -> None:
    """vF 본문쪽이 레이아웃 위에 얹는 것 — 장 번호 배지 · 로고 · 바닥 칩."""
    ph = slide.shapes.title
    ph.text_frame.text = title
    r = ph.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(25)
    _typeface(r, FONT_B)

    box(slide, 11.973, 0.0, 1.319, 0.648, fill=PAPER)
    lg = LOGO / "s4_8.png"
    if lg.exists():
        slide.shapes.add_picture(str(lg), Emu(emu(11.395)), Emu(emu(0.01)),
                                 width=Emu(emu(1.914)), height=Emu(emu(0.408)))
    t = text(slide, 0.045, 0.184, 0.693, 0.509, [(tag, 22, True, WHITE)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    t.text_frame.word_wrap = False

    box(slide, 0.0, 7.28, 2.989, 0.197, fill=RULE)
    text(slide, 0.12, 7.28, 2.8, 0.197, [("2026 KAIA 자문회의", 12, True, SLATE)],
         anchor=MSO_ANCHOR.MIDDLE)


# ── ③ 표지 ────────────────────────────────────────────────────────────────
def cover(slide) -> None:
    box(slide, 0, 0.55, SW, 2.30, fill=NAVY)
    box(slide, 0.74, 0.0, 0.28, 2.85, fill=WHITE)
    text(slide, 0.75, 1.00, 11.96, 1.30,
         [[("위성 InSAR · PINN 기반 교량 변위 모니터링 플랫폼", 30, True, WHITE)],
          [("Inframon", 30, True, WHITE)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.12)
    text(slide, 0.75, 2.98, 11.96, 0.30,
         [("좌표 하나로 위성 → 변위 → AI 해석 → 디지털 트윈 → B-Maps 까지 자동으로 잇는다",
           13.5, True, SLATE)], align=PP_ALIGN.CENTER)
    text(slide, 0.75, 3.30, 11.96, 0.26,
         [("KICT B-Maps 연계 · 2026 KAIA 자문회의 · 강원대학교", 11, False, GRAY)],
         align=PP_ALIGN.CENTER)

    y = 3.78
    for num, ti, sub, pg in CONTENTS:
        box(slide, 1.55, y, 0.58, 0.58, fill=BLUE)
        text(slide, 1.55, y, 0.58, 0.58, [(num, 15, True, WHITE)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 2.35, y + 0.02, 8.6, 0.28, [(ti, 15, True, SLATE)])
        text(slide, 2.35, y + 0.31, 8.9, 0.24, [(sub, 10.5, False, GRAY)])
        text(slide, 11.15, y + 0.02, 1.2, 0.28, [(pg, 13, True, BLUE)],
             align=PP_ALIGN.RIGHT)
        box(slide, 1.55, y + 0.62, 10.8, 0.006, fill=RULE)
        y += 0.68

    for name, x, yy, w, h in COVER_LOGOS:
        p = LOGO / name
        if p.exists():
            slide.shapes.add_picture(str(p), Emu(emu(x)), Emu(emu(yy)),
                                     width=Emu(emu(w)), height=Emu(emu(h)))
    box(slide, 0, 7.17, SW, 0.29, fill=PAPER)
    text(slide, 0, 7.17, SW, 0.29, [(AGENCIES, 9.5, False, SLATE)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--vf", default=str(VF))
    ap.add_argument("--template", default=str(TEMPLATE))
    ap.add_argument("--src", default=str(SRC))
    ap.add_argument("--out", default=str(OUT))
    ap.add_argument("--make-template", action="store_true",
                    help="vF 에서 틀을 새로 뽑는다(vF 파일이 있어야 한다)")
    a = ap.parse_args()

    if a.make_template or not Path(a.template).exists():
        if not Path(a.vf).exists():
            print(f"!! vF 자료가 없다: {a.vf}")
            return 1
        make_template(Path(a.vf), Path(a.template))

    prs = Presentation(a.template)
    src = Presentation(a.src)
    lays = {ly.name: ly for ly in prs.slide_masters[0].slide_layouts}
    page = lays["page 1"]

    s0 = prs.slides.add_slide(lays["Title and Content"])
    for sh in list(s0.shapes):          # 빈 자리표시자를 치운다
        sh._element.getparent().remove(sh._element)
    cover(s0)
    print("  1쪽  표지 + 목차 (vF 표지 형식)")

    for n, (sn, tag, title) in enumerate(PAGES, start=2):
        s = prs.slides.add_slide(page)
        chrome(s, tag, title)
        ssrc = src.slides[sn - 1]
        kept = 0
        for i, sh in enumerate(ssrc.shapes):
            if i in CHROME:
                continue
            copy_shape(s, ssrc, sh, 2000 + n * 200 + i * 4)
            kept += 1
        print(f"  {n}쪽  {tag} {title[:34]:<36} 도형 {kept}개")

    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.out)
    print(f"만들었다: {a.out}  ({os.path.getsize(a.out) / 1e6:.2f} MB · "
          f"{len(prs.slides.__iter__.__self__._sldIdLst)}쪽)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
