#!/usr/bin/env python3
"""7p 자료 — **성능은 두루뭉실하게, 무게는 앞으로의 기대에 싣는다**.

`patch_deck_drop_gnss.py` 가 계측 대조를 걷어냈지만 성능처럼 읽히는 말이 아직 남아
있었다. 지금 단계에서 숫자로 다투면 지는 싸움이다 — 교면 측점이 아직 얇고, 속도의
95 % 구간은 대개 0 을 품는다. 그러니 **무엇이 되는지**와 **앞으로 무엇을 채우는지**로
말을 옮긴다.

남아 있던 것
  · 슬라이드 4 — "품질 게이트 / 기준 미달 결과는 자동으로 **전송 보류**"
    (게이트 낱말은 이미 다른 자리에서 다 걷어냈는데 여기만 남아 있었다)
    "**16/16** 개소" 전수 비율, "전 과정 **관통**" 같은 완료형
  · 슬라이드 5 — 표의 **연 변위속도** 열. 교량마다 +0.17 ~ −4.26 mm/년 이 찍혀 있는데
    대부분 95 % 구간이 0 을 품는다. 점값만 실으면 확정된 관측처럼 읽힌다
  · 슬라이드 3 — 표 머리 "**현재 달성**"
  · 슬라이드 7 — "**개발은 끝났다**", "23개소 관통"
  · 슬라이드 3·4 — "**공명** 위험 지수", "잔존수명 추정"
    (자문회의에서 '공명' 이 어렵다는 말이 나와 v2 자료에서 이미 풀어 썼다)

    python scripts/patch_deck_soften_7p.py [--dry-run]

산출: 대상 pptx 를 제자리에서 갱신 · docs/bridges/deck_soften_7p_log.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from patch_deck_drop_gnss import drop_col                             # noqa: E402
from patch_kaia_v2_slide13 import set_para, style_groups               # noqa: E402

DECK = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_1.pptx"
LOG: list = []

# (슬라이드(1부터), 도형, 문단, 바꾸기 전 글의 일부, [강조 묶음별 글], 이유)
EDITS = [
    # ── Ⅱ 연구성과 목표 ───────────────────────────────────────────────────
    (3, 32, 1, "공명 위험 지수(CRI) 산출",
     ["위험 지수(CRI) 산출"],
     "'공명' 은 물리 공진으로 오독된다 — 자문회의에서 어렵다는 말이 나왔다"),
    (3, 44, 0, "※ 현재 달성 = 실제 Sentinel-1",
     ["※ 지금 단계 = 실제 Sentinel-1 자료(ASC path127, 2018-06-19 ~ 2025-12-27)로 "
      "전 과정을 돌려 본 결과다. 정확도 지표는 계측 대조와 측점 밀도 확보 뒤에 제시한다."],
     "각주에도 '달성' 대신 '지금 단계' 로, 정확도는 앞으로 낸다고 적는다"),
    # ── Ⅲ① 플랫폼 구성 ───────────────────────────────────────────────────
    (4, 46, 0, "④ FRAM · 위험도", ["④ 위험도 평가"],
     "제목에서 어려운 말을 뺀다"),
    (4, 47, 0, "공명 위험 지수 CRI", ["위험 지수 CRI 산출"],
     "'공명' 을 푼다"),
    (4, 47, 2, "잔존수명 추정", ["등급 기준은 잠정"],
     "'잔존수명 추정' 은 지금 자료로 뒷받침되지 않는다 — 잠정임을 먼저 적는다"),
    (4, 60, 0, "품질 게이트", ["근거 기록"],
     "게이트 낱말은 구조 등급으로 오독된다"),
    (4, 60, 1, "전송 보류",
     ["관측 조건과 신뢰구간을 결과에 함께 남긴다"],
     "'기준 미달 결과는 자동으로 전송 보류' — 다른 자리에서 다 걷어낸 말이 여기만 남아 있었다"),
    (4, 70, 0, "전 과정 관통 교량", ["전 과정 시범 처리 교량"],
     "'관통' 은 완료형이다"),
    (4, 72, 0, "16/16", ["16", "  개소"],
     "전수 비율은 검증 성적처럼 읽힌다 — 개소 수만"),
    (4, 73, 0, "한강 교량 IFC 트윈", ["한강 교량 트윈 시범"],
     "시범임을 적는다"),
    # ── Ⅲ② 한강 교량 실증 ────────────────────────────────────────────────
    (5, 20, 0, "연 변위속도 = 교면 측점 중앙값",
     ["관측 시점 수와 교면 측점 수는 교량마다 다르다 — 변위속도와 정밀도는 "
      "측점 밀도를 확보한 뒤에 제시한다"],
     "속도 열을 뺐으므로 설명도 '앞으로 낸다' 로"),
    # ── Ⅳ 향후 계획 ──────────────────────────────────────────────────────
    (7, 14, 0, "개발은 끝났다",
     ["파이프라인은 끝까지 돌아간다 — KICT 와 연동 규칙을 정하면 시범 연동부터 "
      "시작할 수 있다"],
     "'개발은 끝났다' 는 더 물을 것이 없다는 말이 된다 — 돌아간다는 사실만"),
    (7, 23, 1, "23개소 관통",
     ["실 Sentinel-1 자료로 전 과정 시범 처리"],
     "'관통' 과 개소 수를 뺀다"),
    (7, 23, 2, "REST 12개", ["B-Maps REST 사이드카 · 탭 예제"],
     "개수보다 무엇인지가 먼저다"),
    (7, 21, 0, "완료", ["진행"],
     "머리글을 '끝났다' 에서 '돌아간다' 로 바꿔 놓고 배지만 '완료' 면 어긋난다 — "
     "배지 폭이 0.85in 이라 같은 두 글자로 맞춘다"),
]

# 표 칸 하나만 바꾸는 것들: (슬라이드, 행, 열, 바꾸기 전, 바꾼 뒤, 이유)
CELLS = [
    (3, 0, 2, "현재 달성", "지금 단계", "'달성' 은 끝났다는 말이다"),
    (3, 4, 2, "REST 12개 · 탭 예제 구현", "REST 사이드카 · 탭 예제 구현",
     "개수보다 무엇인지가 먼저다"),
]


def cell_text(cell, text) -> str:
    pa = cell.text_frame.paragraphs[0]
    old = "".join(r.text for r in pa.runs)
    if pa.runs:
        pa.runs[0].text = text
        for r in pa.runs[1:]:
            r.text = ""
    return old


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--log", default="docs/bridges/deck_soften_7p_log.json")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    ok = True

    for sn, shi, pi, needle, parts, why in EDITS:
        pa = prs.slides[sn - 1].shapes[shi].text_frame.paragraphs[pi]
        cur = "".join(r.text for r in pa.runs)
        if needle not in cur:
            LOG.append({"찾음": False, "쪽": sn, "도형": shi, "문단": pi,
                        "찾던 글": needle, "그 자리 글": cur[:90]})
            ok = False
            continue
        g = style_groups(pa.runs)
        before = [(None if pa.runs[i[0]].font.size is None
                   else pa.runs[i[0]].font.size.pt, pa.runs[i[0]].font.bold) for i in g]
        old = cur if a.dry_run else set_para(prs.slides[sn - 1].shapes[shi], pi, parts)
        LOG.append({"쪽": sn, "도형": shi, "문단": pi, "이전": old,
                    "이후": "".join(parts), "이유": why, "서식(크기pt,볼드)": before})

    for sn, ri, ci, needle, new, why in CELLS:
        tb = next(s for s in prs.slides[sn - 1].shapes if s.has_table).table
        cur = tb.cell(ri, ci).text
        if needle not in cur:
            LOG.append({"찾음": False, "쪽": sn, "표칸": f"r{ri}c{ci}",
                        "찾던 글": needle, "그 자리 글": cur[:90]})
            ok = False
            continue
        old = cur if a.dry_run else cell_text(tb.cell(ri, ci), new)
        LOG.append({"쪽": sn, "표칸": f"r{ri}c{ci}", "이전": old, "이후": new,
                    "이유": why})

    # 슬라이드 5 — 연 변위속도 열을 통째로 뺀다.
    tb5 = next(s for s in prs.slides[4].shapes if s.has_table).table
    if tb5.cell(0, 4).text.strip() == "연 변위속도":
        if not a.dry_run:
            drop_col(tb5, 4,
                     "점값만 실으면 확정된 관측처럼 읽힌다 — 대부분 95 % 구간이 0 을 품는다")
        LOG.append({"쪽": 5, "표 열 삭제": "연 변위속도",
                    "이유": "점값만 실으면 확정된 관측처럼 읽힌다 — "
                          "대부분 95 % 구간이 0 을 품는다"})
    else:
        LOG.append({"찾음": False, "쪽": 5, "찾던 글": "연 변위속도 열",
                    "그 자리 글": tb5.cell(0, 4).text[:40]})
        ok = False

    for e in LOG:
        if e.get("찾음") is False:
            print(f"  ⚠ {e['쪽']}쪽 못 찾음: {e['찾던 글']} / 그 자리: {e.get('그 자리 글')}")
        elif "표 열 삭제" in e:
            print(f"  − {e['쪽']}쪽 표 열 삭제: {e['표 열 삭제']}")
        else:
            print(f"  · {e['쪽']}쪽 {e['이전'][:50]}")
            print(f"       → {e['이후'][:50]}")
    if not ok:
        print("!! 자리가 안 맞는다 — 아무것도 저장하지 않는다")
        return 1
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0

    prs.save(a.deck)
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name,
         "_왜": "지금 단계에서 숫자로 다투면 지는 싸움이다 — 교면 측점이 아직 얇고 속도의 "
              "95 % 구간은 대개 0 을 품는다. 무엇이 되는지와 앞으로 무엇을 채우는지로 옮긴다.",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {a.deck}")
    print("wrote", a.log)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
