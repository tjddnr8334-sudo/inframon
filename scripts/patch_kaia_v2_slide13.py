#!/usr/bin/env python3
"""`2026 KAIA 자문회의 ppt_v2.pptx` 13쪽 — **어려운 말을 풀고, 다 됐다는 말을 뺀다**.

자문회의에서 두 가지가 걸렸다.
  ① "FRAM 공명" 이 어렵다. 공명은 물리 공진이 아니라 기능공명분석(FRAM)의 비유인데,
     구조 하는 사람이 들으면 고유진동수 이야기로 읽힌다. 말을 풀어 쓰고 학술 용어는
     괄호로 내린다.
  ② "실증 데이터 검증 완료 (한강 교량 16개소)" — 검증이 끝났다고 읽힌다. 끝나지
     않았다. 계측 대조는 지금 값으로 아무것도 주장하지 못한다
     (docs/bridges/strip_log.json). **앞으로 하겠다**로 바꾼다.

②의 취지를 같은 쪽 다른 자리에도 맞춘다 — 한 쪽 안에서 어떤 칸은 "완료", 어떤 칸은
"향후" 면 그게 더 애매하다.
  · "순수 유효 산란체(103점) 100% 분리" · "103점 100% 매핑, 전체 94%" — 비율은
    검증 성적처럼 읽힌다
  · "'보고 가능/조건부/불가' 및 잠정 경보를 스스로 통제" — 게이트 낱말은 구조 등급으로
    오독된다(교량 결과물·7p 자료에서도 같은 이유로 걷어냈다)

대상 파일은 250 MB 제3자 자료라 리포에 담지 않는다(.gitignore). 이 스크립트만 남겨
무엇을 왜 바꿨는지 따라갈 수 있게 한다.

    python scripts/patch_kaia_v2_slide13.py

산출: 대상 pptx 를 제자리에서 갱신 · docs/bridges/kaia_v2_slide13_log.json
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "2026 KAIA 자문회의 ppt_v2.pptx"
LOG: list = []

# (도형 번호, 문단 번호, 바꾸기 전 글의 일부, [강조 묶음별 글], 이유)
#
# 글 목록은 **원문의 볼드 묶음 수와 뜻에 맞춘다**(아래 set_para 참고). 이 자료의 본문
# 문단은 「볼드 머리말 : 일반 설명 + 끝에 볼드 강조」 꼴이라, 한 덩어리로 넣으면
# 크기는 지켜져도 강조가 통째로 사라진다. 강조 자리에는 그 문장에서 **정말 남길 말**을
# 넣는다 — 전에는 '100% 분리' 같은 성적이 강조돼 있었다.
EDITS = [
    (7, 9, "FRAM 공명 위험 지수 산출",
     ["위험 지수(CRI) 산출 (Risk Evaluation)"],
     "'FRAM 공명' 은 물리 공진으로 오독된다 — 목록에서는 이름만 남기고 설명은 기능 03 으로"),
    (7, 15, "실증 데이터 검증 완료",
     ["한강 교량 시범 적용 (계측 대조 검증은 향후 과제)"],
     "검증이 끝났다고 읽힌다 — 끝나지 않았다. 앞으로 하겠다로"),
    (16, 4, "100% 분리",
     ["성수대교 시범: ",
      "OSM 데크선 쉬프트 보정과 DEM 잔차고도 분석으로 교면 위 측점을 가려내는 절차를 "
      "적용(103점). ",
      "측점 밀도 확보는 향후 과제"],
     "'100% 분리' 는 비율이라 검증 성적처럼 읽힌다 — 절차를 적용했다로 바꾸고, "
     "강조는 남은 과제에 둔다"),
    (17, 6, "기능 03. FRAM 공명 위험 지수(CRI)",
     ["기능 03. 위험 지수(CRI) 평가 & 자가 점검"],
     "제목에서 어려운 말을 뺀다"),
    (17, 8, "기능공명분석(FRAM) 기반 기능별 변동",
     ["CRI 지수 도출: ",
      "여러 지표의 흔들림이 서로 맞물려 커지는 정도를 0~1 지수로 환산 — 하나가 기준을 "
      "넘었는지가 아니라 ",
      "여럿이 겹치는지를 본다 (기능공명분석 FRAM 기반)"],
     "무엇을 보는 지수인지 평이한 말로 먼저 적고, 학술 용어는 괄호로 내린다"),
    (17, 10, "'보고 가능/조건부/불가'",
     ["신뢰성 자가 점검: ",
      "관측 잡음(σ)과 95% 신뢰구간(CI)을 매번 함께 계산해, ",
      "지금 자료로 말할 수 있는 범위를 결과에 같이 적는다"],
     "'보고 가능/조건부/불가' 는 구조 등급으로 오독된다 — 하는 일만 적는다"),
    (18, 2, "103점 100% 매핑, 전체 94%",
     ["부재 결합: ",
      "실측 제원 기반 IFC4 프록시 구조물의 부재 고유식별자(GUID)에 변위 시계열을 "
      "결합 (성수대교 ",
      "103점",
      ")"],
     "비율은 검증 성적처럼 읽힌다 — 무엇을 결합하는지만"),
]


def _style(r):
    """그 run 의 겉모습 — 볼드와 크기. 이 둘이 같으면 한 묶음으로 본다.

    처음에는 볼드만 봤는데, 큰 숫자 + 작은 단위로 된 타일("16/16" 24pt + "개소" 12pt)
    이 한 묶음으로 뭉쳐 버렸다. 크기도 함께 봐야 원문의 생김새가 지켜진다.
    """
    return (r.font.bold, None if r.font.size is None else r.font.size.pt)


def style_groups(runs) -> list:
    """겉모습이 같은 이웃 run 끼리 묶는다 — 그 묶음이 원문의 강조 구조다."""
    groups: list = []
    for i, r in enumerate(runs):
        if groups and _style(runs[groups[-1][0]]) == _style(r):
            groups[-1].append(i)
        else:
            groups.append([i])
    return groups


def set_para(shape, idx: int, parts: list) -> str:
    """문단의 **글자만** 바꾼다 — 있는 run 에 나눠 넣고 서식은 한 군데도 손대지 않는다.

    run 을 지우고 새로 달면 파일이 깨져 PowerPoint 가 못 연다(전에 겪었다). 그래서
    `run.text` 만 바꾸는데, 그것만으로는 부족하다. 첫 run 에 전부 몰아넣으면 크기는
    지켜지지만 **원문의 강조가 통째로 사라진다** — 이 자료는 「볼드 머리말 : 일반 설명
    + 끝에 볼드 강조」 꼴이기 때문이다.

    그래서 겉모습(볼드·크기)이 같은 이웃 run 을 묶고, `parts[i]` 를 i번째 묶음의 첫 run 에 넣는다.
    나머지 run 은 빈 글자로 둔다(지우지 않는다). 크기·글꼴·색·볼드는 전부 있던 run 의
    것을 그대로 쓰므로, 바뀌는 것은 글자뿐이다.
    """
    pa = shape.text_frame.paragraphs[idx]
    runs = pa.runs
    if not runs:
        raise ValueError(f"문단 {idx} 에 run 이 없다")
    old = "".join(r.text for r in runs)

    groups = style_groups(runs)
    if len(parts) > len(groups):
        raise ValueError(f"문단 {idx}: 글 {len(parts)}조각인데 강조 묶음은 "
                         f"{len(groups)}개뿐이다 — 넣을 자리가 없다")
    texts = [""] * len(runs)
    for gi, part in enumerate(parts):
        texts[groups[gi][0]] = part
    for r, t in zip(runs, texts):
        r.text = t
    return old


def swap_part(src, dest, partname, blob) -> None:
    """**그 슬라이드 파트 하나만** 갈아 끼우고 나머지는 바이트 그대로 복사한다.

    `Presentation.save()` 로 통째로 다시 쓰면 안 된다. 실제로 그렇게 했더니 86개 파트가
    다시 쓰였고([Content_Types].xml · slide18.xml.rels 포함) **PowerPoint 가 파일을
    열지 못했다.** zip 도 성하고 python-pptx 는 다시 읽는데 PowerPoint 만 거부한다 —
    이 자료에는 python-pptx 가 온전히 되돌리지 못하는 부분이 있다는 뜻이다.

    그래서 원본 zip 을 항목 단위로 베끼되 대상 파트만 새 XML 로 바꾼다. 압축방식·
    타임스탬프·속성을 그대로 들고 가므로, 바뀌는 것은 그 슬라이드의 글자뿐이다.
    """
    import shutil
    import tempfile
    import zipfile

    want = str(partname).lstrip("/")
    tmp = Path(tempfile.mkdtemp()) / "out.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w") as zout:
        if want not in zin.namelist():
            raise KeyError(f"{want} 가 원본에 없다")
        for item in zin.infolist():
            data = blob if item.filename == want else zin.read(item.filename)
            zout.writestr(item, data)          # 압축방식·날짜·속성을 그대로
    shutil.move(str(tmp), str(dest))
    shutil.rmtree(tmp.parent, ignore_errors=True)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(DECK))
    ap.add_argument("--slide", type=int, default=13)
    ap.add_argument("--log", default="docs/bridges/kaia_v2_slide13_log.json")
    ap.add_argument("--out", default=None,
                    help="다른 곳에 저장한다(원본이 PowerPoint 에 열려 있어 잠겼을 때)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    s = prs.slides[a.slide - 1]
    ok = True
    def shape_of(g, runs):
        """강조 묶음을 (크기pt, 볼드) 목록으로 — 서식이 안 변했음을 눈으로 본다."""
        return [(None if runs[i[0]].font.size is None else runs[i[0]].font.size.pt,
                 runs[i[0]].font.bold) for i in g]

    for shi, pi, needle, parts, why in EDITS:
        sh = s.shapes[shi]
        pa = sh.text_frame.paragraphs[pi]
        cur = "".join(r.text for r in pa.runs)
        if needle not in cur:
            # 자리가 밀렸으면 엉뚱한 문단을 덮어쓴다 — 멈추고 사람이 본다.
            LOG.append({"찾음": False, "도형": shi, "문단": pi,
                        "찾던 글": needle, "그 자리 글": cur[:90]})
            ok = False
            continue
        g = style_groups(pa.runs)
        before = shape_of(g, pa.runs)
        old = cur if a.dry_run else set_para(sh, pi, parts)
        after = shape_of(style_groups(pa.runs), pa.runs)
        LOG.append({"도형": shi, "문단": pi, "이전": old,
                    "이후": "".join(parts), "이유": why,
                    "강조 묶음": f"{len(parts)}/{len(g)} 사용",
                    "서식(크기pt,볼드)": before,
                    "서식 그대로인가": before == after})

    for e in LOG:
        if e.get("찾음") is False:
            print(f"  ⚠ [{e['도형']}]p{e['문단']} 못 찾음: {e['찾던 글']}")
            print(f"       그 자리에는: {e['그 자리 글']}")
        else:
            print(f"  · [{e['도형']}]p{e['문단']} {e['이전'][:56]}")
            print(f"       → {e['이후'][:56]}")
            print(f"       서식 {e['서식(크기pt,볼드)']} · 묶음 {e['강조 묶음']} · "
                  f"그대로인가={e['서식 그대로인가']}")
    if not ok:
        print("!! 문단 자리가 안 맞는다 — 아무것도 저장하지 않는다")
        return 1
    if a.dry_run:
        print("(dry-run — 저장하지 않았다)")
        return 0

    dest = a.out or a.deck
    try:
        swap_part(a.deck, dest, s.part.partname, s.part.blob)
    except PermissionError:
        print(f"!! 저장이 막혔다 — PowerPoint 에서 열려 있는 듯하다: {dest}")
        print("   닫은 뒤 다시 돌리거나, --out 으로 다른 곳에 저장한다.")
        return 1
    Path(a.log).write_text(json.dumps(
        {"_대상": Path(a.deck).name + " (250 MB · 리포에 담지 않는다)",
         "_슬라이드": a.slide,
         "_왜": "① 'FRAM 공명' 이 어렵다 — 물리 공진으로 오독된다. ② '검증 완료' 는 "
              "끝났다고 읽히는데 끝나지 않았다. 애매한 표현을 없애고 향후 과제로 적는다.",
         "바꾼 것": LOG}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고쳤다: {dest}")
    print("wrote", a.log)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
