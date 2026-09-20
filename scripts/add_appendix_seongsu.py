#!/usr/bin/env python3
"""**부록 — 성수대교 산출 결과** 한 쪽을 붙인다.

왜 성수대교인가. 6쪽 B-Maps 동작 화면이 성수대교라 이야기가 이어지고, 관측이
101회로 가장 길며, 무엇보다 **열팽창으로 교면과 지반이 갈라진** 교량이다
(교면 +0.15 ± 0.06 vs 지반 +0.41 ± 0.09 mm/°C · `thermal_separates: true`).
"교면을 보고 있는 게 맞느냐" 는 이 과제에서 가장 아픈 질문인데, 여기서는 그 답이
숫자로 나온다.

쓰지 않는 것: `brief.png`. 그 그림은 **옛 ±30 m 선별** 결과라 '유효 PS 1점 ·
잔차고도 −18.3 m(지면보다 낮음)' 로 나온다 — 우리가 폐기한 분석이다. 대신 교면 전용
재선별 그림(`교면전용.png`)을 쓴다.

숫자는 두 갈래를 **섞지 않고 나눠** 적는다. MT-InSAR 전체 처리와 교면 전용 재도출은
서로 다른 단계이고, 점 수·속도가 다르다(536→204점 vs 15점). 한 칸에 섞으면
"왜 숫자가 다르냐" 가 나온다.

    python scripts/add_appendix_seongsu.py --deck <vF 형식 pptx> [--dry-run]

산출: 대상 pptx 끝에 부록 한 쪽 추가
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from kaia_theme import (                                              # noqa: E402
    BLUE, GRAY, GREEN, INK, PAPER, SLATE, WHITE, box, panel, picture_fit, text,
)
from rebuild_7p_vF import chrome                                      # noqa: E402

VF = ROOT / "docs" / "KICT_Bmaps_Inframon_연구개발_7p_vF.pptx"
BR = ROOT / "docs" / "bridges" / "성수대교"
TITLE = "부록 — 성수대교 산출 결과"


def rows(mt: dict, rb: dict) -> tuple[list, list]:
    """왼쪽(처리) · 오른쪽(교면 결과). 두 단계를 섞지 않는다."""
    net, bl = mt["network"], mt["baseline"]
    left = [
        ("관측", f"{mt['n_epochs']}회 · 2018-06 ~ 2025-12"),
        ("기선망", f"쌍 {net['n_pairs']}개 · 연결됨 · "
                 f"수직기선 {bl['bperp_min_m']:+.0f} ~ {bl['bperp_max_m']:+.0f} m"),
        ("측점", f"{mt['n_points']}점 중 γ≥{mt['gamma_min']:.2f} 를 넘은 "
                f"{mt['n_kept']}점 · 잔차고도 σΔh {mt['sigma_dh_median_m']:.1f} m"),
        ("APS 제거", f"RMS {mt['aps_rms_mm']:.1f} mm · 잔차 "
                   f"{mt['resid_rms_before_mm']:.1f} → {mt['resid_rms_after_mm']:.1f} mm"),
    ]
    right = [
        ("교면 측점", f"{rb['n_deck_ps']}점 · 시간결맞음 γ 중앙 "
                   f"{rb['gamma_median_deck']:.3f} · 코히런스 {rb['coh_mean']:.3f}"),
        ("열팽창", f"교면 {mt['thermal_deck_mm_per_C']:+.2f} vs 지반 "
                f"{mt['thermal_ground_mm_per_C']:+.2f} mm/°C — 갈라진다"),
        ("속도", f"중앙값 {rb['velocity_median_mm_yr']:+.2f} mm/년"),
        ("위험 지수", f"측점 중앙 {rb['cri_point_median']:.3f} · 전역 {rb['cri']:.2f} "
                  f"(등급 기준은 잠정) · IFC 부재에 {rb['bound']}점 결합"),
    ]
    return left, right


def block(slide, x, y, w, h, title, items, *, mark=BLUE):
    panel(slide, x, y, w, h, title, fill=WHITE, tab=SLATE)
    yy = y + 0.46
    for head, body in items:
        box(slide, x + 0.18, yy + 0.05, 0.045, 0.18, fill=mark)
        text(slide, x + 0.32, yy - 0.02, 1.25, 0.22, [(head, 10, True, INK)])
        text(slide, x + 1.60, yy - 0.02, w - 1.82, 0.40, [(body, 9.4, False, SLATE)])
        yy += 0.42


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(VF))
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    mt = json.loads((BR / "mtinsar.json").read_text(encoding="utf-8"))
    rb = json.loads((BR / "재도출.json").read_text(encoding="utf-8"))
    left, right = rows(mt, rb)
    if a.dry_run:
        for nm, items in (("처리(MT-InSAR)", left), ("교면 결과", right)):
            print(f"  ── {nm}")
            for h_, b_ in items:
                print(f"     {h_:<10} {b_}")
        print("(dry-run — 저장하지 않았다)")
        return 0

    prs = Presentation(a.deck)
    lays = {ly.name: ly for ly in prs.slide_masters[0].slide_layouts}
    if "page 1" not in lays:
        print("!! 'page 1' 레이아웃이 없다 — vF 형식 자료에만 붙일 수 있다")
        return 1
    if any(TITLE in (sh.text_frame.text if sh.has_text_frame else "")
           for sl in prs.slides for sh in sl.shapes):
        print("이미 부록이 있다 — 아무것도 하지 않는다")
        return 0

    s = prs.slides.add_slide(lays["page 1"])
    chrome(s, "부록", TITLE)

    box(s, 0.33, 0.83, 0.06, 0.3, fill=BLUE)
    text(s, 0.48, 0.79, 12.4, 0.38,
         [("관측 101회 · 교면 측점 15점 — 열팽창으로 교면과 지반이 갈라진 사례", 14.5,
           True, INK)])

    block(s, 0.45, 1.40, 6.1, 2.26, "처리 — MT-InSAR", left)
    block(s, 6.78, 1.40, 6.1, 2.26, "교면 결과", right, mark=GREEN)

    fig = BR / "교면전용.png"
    if fig.exists():
        # 그림이 가로로 매우 길다(약 3.4:1). 폭을 다 쓰면 세로가 모자라므로 높이에
        # 맞춰 줄이고, **칸이 그림을 감싸게** 한다 — 넓은 회색 판 가운데 작게 박힌
        # 그림은 자리를 못 잡은 것처럼 보인다.
        from PIL import Image
        with Image.open(fig) as im:
            iw, ih = im.size
        avail_h = 2.40
        fw = min(12.09, avail_h * iw / ih)
        fh = fw * ih / iw
        px = (13.333 - (fw + 0.34)) / 2
        panel(s, px, 3.80, fw + 0.34, fh + 0.62, "교면 전용 재선별",
              fill=PAPER, tab=SLATE)
        picture_fit(s, fig, px + 0.17, 4.23, fw, fh)
    text(s, 0.45, 6.88, 12.43, 0.3,
         [("※ 처리(MT-InSAR)와 교면 전용 재선별은 서로 다른 단계라 점 수·속도가 다르다. "
           "옛 ±30 m 선별로 만든 브리프 그림은 쓰지 않는다 — 그 점들은 지면보다 낮게 "
           "나와 교면이 아니었다.", 8.6, False, GRAY)])

    prs.save(a.deck)
    print(f"부록을 붙였다: {a.deck}  (이제 {len(prs.slides.__iter__.__self__._sldIdLst)}쪽)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
